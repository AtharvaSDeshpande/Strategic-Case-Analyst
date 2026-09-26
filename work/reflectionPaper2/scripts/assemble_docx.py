import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Set standard styles
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Title
heading = doc.add_heading('Strategic Reflections: Architecture, Erosion, and Scale', 0)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

base_dir = r"c:\Users\athar\Desktop\Antigravity Playground\Strategic-Case-Analyst\work\reflectionPaper2"
diag_dir = os.path.join(base_dir, "diagrams")

def add_paragraph(text, justify=True):
    p = doc.add_paragraph(text)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def add_figure(filename, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    img_path = os.path.join(diag_dir, filename)
    if os.path.exists(img_path):
        r.add_picture(img_path, width=Inches(6.0))
    else:
        print(f"Warning: {img_path} not found.")
    
    caption = doc.add_paragraph(caption_text)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True

add_paragraph("When we observe how different organizations compete, the most enduring advantages rarely stem from a single brilliant product. Instead, they arise from how a company orchestrates its entire system of activities, manages the inevitable decay of its own innovations, and wrangles complexity as it scales. By examining three very different entities—IKEA in retail, NVIDIA in technology, and Unilever in fast-moving consumer goods—I have come to see that strategic durability requires a constant, often uncomfortable alignment of operational reality with market gravity.")

add_paragraph("For IKEA, the central reality is radical cost engineering (Figure 1). Their entire organizational architecture is designed to make functional design accessible to the masses. This isn’t simply a matter of finding cheaper suppliers; it is a foundational philosophy that dictates every downstream decision.")
add_figure("figure_1_ikea_capability.png", "Figure 1: IKEA Organizational Capability Distribution")

add_paragraph("IKEA’s most recognizable innovation, the flat-pack, perfectly illustrates how they have reimagined the traditional sequence of business activities (Figure 2). Instead of designing a beautiful chair and then figuring out how to ship it, IKEA designers work backward from a strict target price. The constraint forces them to design furniture that eliminates \"dead air\" in packaging. Furthermore, they offload two of the most expensive steps in retail—last-mile delivery and final assembly—directly onto the customer. By requiring shoppers to pull boxes from warehouse shelves and build the furniture at home, IKEA functionally transforms its customers into uncompensated final-stage workers, completely reconfiguring the cost structure of the industry in a way traditional competitors struggle to replicate.")
add_figure("figure_2_ikea_value_chain.png", "Figure 2: IKEA Operating Model Reconfiguration")

add_paragraph("To protect this low-price guarantee, IKEA has had to draw a very specific boundary around what it owns and controls (Figure 3). Rather than expanding into unrelated retail ventures, they have historically reached deep into their own supply chain. By purchasing forests in Eastern Europe and controlling manufacturing facilities, they insulate themselves against raw material price shocks. They don't just sell furniture; they orchestrate the entire material journey from the forest to the living room, ensuring that margin isn't lost to intermediaries along the way.")
add_figure("figure_3_ikea_scope.png", "Figure 3: Vertical Integration for Cost Control")

add_paragraph("In the technology sector, the battle is fought not over physical logistics, but over developer mindshare. NVIDIA’s dominance is often attributed to the raw speed of its silicon, but its true core capability is the parallel computing software ecosystem it has cultivated for nearly two decades (Figure 4).")
add_figure("figure_4_nvidia_capability.png", "Figure 4: NVIDIA Organizational Capability Distribution")

add_paragraph("The gravity of this ecosystem creates an incredibly hostile environment for any challenger attempting to enter the market (Figure 5). Because major artificial intelligence workflows have been written in NVIDIA’s proprietary software language since 2006, the cost for a research lab or enterprise to switch to a rival chip is astronomical. They wouldn't just be buying new hardware; they would have to invest thousands of engineering hours to rewrite and re-optimize their code. This massive friction locks customers in and neutralizes the threat of cheaper alternatives, granting NVIDIA pricing power that hardware specifications alone could never justify.")
add_figure("figure_5_nvidia_5forces.png", "Figure 5: The Software Ecosystem Moat")

add_paragraph("Yet, this lock-in is not enough to guarantee permanent survival in a sector defined by rapid obsolescence. NVIDIA’s survival mechanism is a relentless willingness to destroy its own revenue streams (Figure 6). By releasing a radically faster chip architecture every two years, they deliberately render their previous generation obsolete. They do not wait for a competitor to disrupt them; they disrupt themselves, ensuring that the only alternative to an aging NVIDIA chip is a new NVIDIA chip. This aggressive cycle of cannibalizing their own products is the engine of their continued dominance.")
add_figure("figure_6_nvidia_creative_destruction.png", "Figure 6: Planned Obsolescence via Rapid Innovation")

add_paragraph("Unilever operates in an entirely different reality, where the challenge is not technological obsolescence, but the sheer, crushing weight of global scale. Their core capability lies in managing an enormous, diverse portfolio of localized and global brands (Figure 7).")
add_figure("figure_7_unilever_capability.png", "Figure 7: Unilever Organizational Capability Distribution")

add_paragraph("Historically, Unilever’s pursuit of scale led to a sprawling collection of roughly 1,600 brands, many of which consumed resources without generating meaningful returns. The strategic correction required a brutal pruning of this portfolio (Figure 8). By starving the slow-growing, low-margin regional labels and reallocating that capital to 400 core global powerhouses, they fundamentally shifted their growth trajectory. It became a continuous exercise in identifying which assets throw off cash and which ones merely drain attention, ensuring that the company’s massive footprint translates into actual profitability.")
add_figure("figure_8_unilever_bcg.png", "Figure 8: Brand Portfolio Pruning")

add_paragraph("However, extracting cash from legacy products is only half the equation; finding new pools of profit requires a deliberate shift in focus (Figure 9). As traditional soap and soup categories faced stagnation, Unilever had to look beyond its historic boundaries. They reorganized into five distinct business groups and aggressively acquired premium wellness and prestige beauty brands. This wasn't simply about selling more volume to existing customers; it was a structural pivot into entirely new, higher-margin categories that are less sensitive to commodity pricing pressures, ensuring they remain relevant as consumer spending habits evolve.")
add_figure("figure_9_unilever_ansoff.png", "Figure 9: Growth via Premium Categories")

add_paragraph("Ultimately, these three organizations demonstrate that a lasting advantage is never static. Whether it is IKEA engineering costs out of every step, NVIDIA aggressively rendering its own technology obsolete, or Unilever ruthlessly pruning its portfolio to fund new ventures, the most successful companies are those that continuously align their internal capabilities with the brutal realities of their respective markets.")

doc.save(os.path.join(base_dir, "Reflection_Paper.docx"))
print("Saved Reflection_Paper.docx")
